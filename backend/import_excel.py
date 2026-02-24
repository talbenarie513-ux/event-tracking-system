import openpyxl
from datetime import datetime, timedelta
from database import Database

def parse_date(date_str):
    """Parse various date formats"""
    if not date_str or date_str == 'None':
        return None
    
    if isinstance(date_str, datetime):
        return date_str.strftime('%Y-%m-%d')
    
    date_str = str(date_str).strip()
    
    if not date_str or date_str == '':
        return None
    
    formats = [
        '%Y-%m-%d %H:%M:%S',
        '%d/%m/%Y',
        '%d.%m.%Y',
        '%d.%m.%y',
        '%Y-%m-%d',
        '%d-%m-%Y'
    ]
    
    for fmt in formats:
        try:
            dt = datetime.strptime(date_str, fmt)
            return dt.strftime('%Y-%m-%d')
        except:
            continue
    
    return None

def map_urgency(hebrew_urgency):
    """Map Hebrew urgency to standard format"""
    if not hebrew_urgency:
        return 'בינונית'
    
    urgency_str = str(hebrew_urgency).strip()
    
    if not urgency_str or urgency_str == '':
        return 'בינונית'
    
    mapping = {
        'גבוה': 'גבוהה',
        'נמוך': 'נמוכה',
        'בינוני': 'בינונית',
        'קריטי': 'קריטית',
        'גבוהה': 'גבוהה',
        'נמוכה': 'נמוכה',
        'בינונית': 'בינונית',
        'קריטית': 'קריטית'
    }
    
    result = mapping.get(urgency_str, 'בינונית')
    
    allowed = ['קריטית', 'גבוהה', 'בינונית', 'נמוכה']
    if result not in allowed:
        return 'בינונית'
    
    return result

def parse_priority(priority_value):
    """Parse priority, handling text values"""
    if not priority_value:
        return 5
    
    try:
        p = int(priority_value)
        if 1 <= p <= 10:
            return p
        return 5
    except:
        pass
    
    return 5

def map_status(status_value):
    """Map various status values to valid database status"""
    if not status_value:
        return 'אירוע חדש'
    
    status_str = str(status_value).strip()
    
    if not status_str or status_str == '':
        return 'אירוע חדש'
    
    valid_statuses = [
        'אירוע חדש', 'בבדיקה', 'בטיפול', 'טופל חלקית',
        'בהכנת הצעת מחיר', 'בפיתוח', 'בבדיקת איכות של פיתוח',
        'ממתין לאישור הצעת מחיר', 'בבדיקת תחום תכנון',
        'הושלם הטיפול', 'בהקפאה'
    ]
    
    if status_str in valid_statuses:
        return status_str
    
    status_mapping = {
        'חדש': 'אירוע חדש',
        'בדיקה': 'בבדיקה',
        'טיפול': 'בטיפול',
        'פיתוח': 'בפיתוח',
        'הושלם': 'הושלם הטיפול',
        'הסתיים': 'הושלם הטיפול',
        'נסגר': 'הושלם הטיפול',
        'סגור': 'הושלם הטיפול',
        'הקפאה': 'בהקפאה',
        'הוקפא': 'בהקפאה',
        'הצעת מחיר': 'בהכנת הצעת מחיר',
        'המתנה': 'ממתין לאישור הצעת מחיר',
        'ממתין': 'ממתין לאישור הצעת מחיר',
        'איכות': 'בבדיקת איכות של פיתוח',
        'תכנון': 'בבדיקת תחום תכנון'
    }
    
    for keyword, valid_status in status_mapping.items():
        if keyword in status_str:
            return valid_status
    
    if len(status_str) > 30:
        if 'טופל' in status_str and 'חלק' in status_str:
            return 'טופל חלקית'
        elif 'טופל' in status_str or 'הושלם' in status_str:
            return 'הושלם הטיפול'
        elif 'פיתוח' in status_str:
            return 'בפיתוח'
        elif 'הקפא' in status_str or 'המתנה' in status_str:
            return 'בהקפאה'
        else:
            return 'בטיפול'
    
    return 'בטיפול'

def clean_text(value):
    """Clean text value"""
    if value is None:
        return ''
    
    text = str(value).strip()
    
    if text == '' or text.lower() == 'none':
        return ''
    
    return text

def has_data_in_row(sheet, row_idx):
    """Check if row has any data in columns B-K"""
    for col_idx in range(2, 12):  # Columns B to K
        cell_value = sheet.cell(row_idx, col_idx).value
        if cell_value and str(cell_value).strip() != '':
            return True
    return False

def import_excel_data(excel_path):
    """Import data from Excel file to database with CORRECT column mapping"""
    db = Database()
    conn = db.get_connection()
    cursor = conn.cursor()
    
    print("=" * 100)
    print("📥 STARTING EXCEL IMPORT")
    print("=" * 100)
    print(f"📁 File: {excel_path}")
    print()
    
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb.active
    
    # Find header row (row with "#" in column A)
    header_row = None
    for row_idx in range(1, 20):
        cell_value = sheet.cell(row_idx, 1).value
        if cell_value and str(cell_value).strip() == '#':
            header_row = row_idx
            break
    
    if not header_row:
        print("❌ Could not find header row with '#' in column A!")
        return
    
    print(f"✅ Found header row at: Row {header_row}")
    print()
    print("=" * 100)
    print("📋 COLUMN MAPPING:")
    print("=" * 100)
    print("Column A (#)        → event_id (auto-generate if empty)")
    print("Column B (תאריך)   → registration_date")
    print("Column C (נושא)    → system → Displays in 'מערכת'")
    print("Column D (לקוח)    → affected_customers → Displays in 'לקוחות מושפעים'")
    print("Column E (מהות)    → event_summary → Displays in 'תמצית האירוע'")
    print("Column F (פניה)    → first_contact_date")
    print("Column G (סטטוס)   → status")
    print("Column H (לוז)     → status_deadline")
    print("Column I (דחיפות)  → urgency")
    print("Column J (עדיפות)  → priority")
    print("Column K (הערות)   → event_details → Displays in 'פירוט האירוע'")
    print("=" * 100)
    print()
    
    # FIRST PASS: Collect all existing event IDs from column A
    existing_ids = set()
    for row_idx in range(header_row + 1, sheet.max_row + 1):
        if not has_data_in_row(sheet, row_idx):
            continue
        
        event_num_cell = sheet.cell(row_idx, 1).value  # Column A
        if event_num_cell:
            try:
                event_id = int(event_num_cell)
                existing_ids.add(event_id)
            except:
                pass
    
    # Determine next auto-generated ID
    if existing_ids:
        next_auto_id = max(existing_ids) + 1
    else:
        next_auto_id = 1
    
    print(f"🔢 Found {len(existing_ids)} rows with explicit IDs")
    print(f"🆕 Auto-generated IDs will start from: {next_auto_id}")
    print(f"🎯 Target: Import exactly 32 events")
    print()
    
    imported_count = 0
    skipped_count = 0
    auto_generated_count = 0
    
    # SECOND PASS: Import all data rows
    for row_idx in range(header_row + 1, sheet.max_row + 1):
        # Check if row has any data
        if not has_data_in_row(sheet, row_idx):
            continue
        
        # Stop after 32 events
        if imported_count >= 32:
            print(f"\n✋ Reached 32 events limit, stopping")
            break
        
        # ===== COLUMN A: Event ID (auto-generate if empty) =====
        event_num_cell = sheet.cell(row_idx, 1).value
        
        if not event_num_cell or str(event_num_cell).strip() == '':
            event_num = next_auto_id
            next_auto_id += 1
            auto_generated_count += 1
            print(f"🆕 Row {row_idx}: Auto-generated ID = {event_num}")
        else:
            try:
                event_num = int(event_num_cell)
            except:
                print(f"⚠️ Row {row_idx}: Invalid ID '{event_num_cell}', skipping")
                skipped_count += 1
                continue
        
        # ===== READ ALL COLUMNS =====
        registration_date = parse_date(sheet.cell(row_idx, 2).value)       # Column B
        system = clean_text(sheet.cell(row_idx, 3).value)                  # Column C → מערכת
        affected_customers = clean_text(sheet.cell(row_idx, 4).value)      # Column D → לקוחות מושפעים
        event_summary = clean_text(sheet.cell(row_idx, 5).value)           # Column E → תמצית האירוע
        first_contact = parse_date(sheet.cell(row_idx, 6).value)           # Column F
        status = map_status(sheet.cell(row_idx, 7).value)                  # Column G
        status_deadline = parse_date(sheet.cell(row_idx, 8).value)         # Column H
        urgency = map_urgency(sheet.cell(row_idx, 9).value)                # Column I
        priority = parse_priority(sheet.cell(row_idx, 10).value)           # Column J
        event_details = clean_text(sheet.cell(row_idx, 11).value)          # Column K → פירוט האירוע
        
        # ===== HANDLE EMPTY VALUES =====
        if not system:
            system = 'לא צוין'
        
        if not affected_customers:
            affected_customers = 'לא צוין'
        
        if not event_summary:
            event_summary = f'אירוע #{event_num}'
        
        if not event_details:
            event_details = ''
        
        # Set default dates if missing
        if not registration_date:
            registration_date = datetime.now().strftime('%Y-%m-%d')
        
        if not status_deadline:
            try:
                reg_date = datetime.strptime(registration_date, '%Y-%m-%d')
                deadline_date = reg_date + timedelta(days=3)
                status_deadline = deadline_date.strftime('%Y-%m-%d')
            except:
                status_deadline = (datetime.now() + timedelta(days=3)).strftime('%Y-%m-%d')
        
        try:
            # Check if event ID already exists
            cursor.execute('SELECT id FROM events WHERE id = ?', (event_num,))
            if cursor.fetchone():
                print(f"⚠️ Event #{event_num} already exists in database, skipping")
                skipped_count += 1
                continue
            
            # ===== INSERT INTO DATABASE =====
            cursor.execute('''
                INSERT INTO events (
                    id, registration_date, first_contact_date, system, event_summary,
                    event_details, affected_customers, urgency, priority,
                    status, status_details, status_deadline, additional_notes,
                    created_by, event_classification
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (
                event_num,                  # Column A (or auto-generated)
                registration_date,          # Column B
                first_contact,              # Column F
                system,                     # Column C → מערכת
                event_summary,              # Column E → תמצית האירוע
                event_details,              # Column K → פירוט האירוע
                affected_customers,         # Column D → לקוחות מושפעים
                urgency,                    # Column I
                priority,                   # Column J
                status,                     # Column G
                '',                         # status_details (empty)
                status_deadline,            # Column H
                '',                         # additional_notes (empty)
                'Import',                   # created_by
                ''                          # event_classification (empty)
            ))
            
            imported_count += 1
            
            # Print verification
            print(f"✅ Event #{event_num} imported (Excel Row {row_idx})")
            print(f"   📅 Date: {registration_date}")
            print(f"   🏢 מערכת (C): {system[:50]}")
            print(f"   📝 תמצית (E): {event_summary[:50]}")
            print(f"   👥 לקוח (D): {affected_customers[:50]}")
            if event_details:
                print(f"   📄 פירוט (K): {event_details[:50]}")
            print()
            
        except Exception as e:
            print(f"❌ Error importing Event #{event_num} (Row {row_idx}): {e}")
            skipped_count += 1
    
    conn.commit()
    conn.close()
    
    print("\n" + "=" * 100)
    print("📊 IMPORT SUMMARY")
    print("=" * 100)
    print(f"✅ Successfully imported: {imported_count} events")
    print(f"🆕 Auto-generated IDs: {auto_generated_count} events")
    if skipped_count > 0:
        print(f"⚠️ Skipped: {skipped_count} rows")
    print()
    
    if imported_count == 32:
        print("🎉 PERFECT! Imported exactly 32 events as expected!")
    else:
        print(f"⚠️ WARNING: Expected 32 events but imported {imported_count}")
    
    print("=" * 100)

if __name__ == '__main__':
    excel_path = r'C:\Users\tal_ba\Documents\DataAnalysis-Tal\מערכת תקלות-פיתוח\חוברת1.xlsx'
    import_excel_data(excel_path)